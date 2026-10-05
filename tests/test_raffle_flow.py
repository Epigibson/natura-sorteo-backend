from datetime import datetime, timedelta, timezone

import pytest
from bson import ObjectId

from app.main import run_maintenance
from app.services import raffle_service as rs

PHONE = "5511223344"


def claim(client, slug, folio, phone=PHONE, name="Ana Pérez"):
    return client.post(f"/api/v1/public/raffles/{slug}/claim", json={"folio": folio, "name": name, "phone": phone})


def pay(client, admin, rid, folio):
    return client.post(f"/api/v1/raffles/{rid}/tickets/{folio}/pay", headers=admin, json={})


# ---------- reclamo ----------
def test_claim_ok_and_taken(client, raffle):
    r = claim(client, raffle["slug"], 1)
    assert r.status_code == 200 and r.json()["code"]
    assert claim(client, raffle["slug"], 1, phone="5599887766").status_code == 400  # ya tomado


def test_claim_respects_limit_and_zero_means_unlimited(client, admin, db, raffle):
    s = raffle["slug"]
    assert claim(client, s, 1).status_code == 200
    assert claim(client, s, 2).status_code == 200
    assert claim(client, s, 3).status_code == 400  # límite 2
    db.raffles.update_one({"slug": s}, {"$set": {"max_tickets_per_person": 0}})
    assert claim(client, s, 3).status_code == 200  # 0 = sin límite (antes bloqueaba a todos)


def test_claim_closed_raffle(client, admin, raffle):
    client.post(f"/api/v1/raffles/{raffle['id']}/close", headers=admin)
    assert claim(client, raffle["slug"], 1).status_code == 400


def test_claim_unpaid_per_ip_cap(client, db, raffle, monkeypatch):
    import app.routers.public as pub
    monkeypatch.setattr(pub, "MAX_UNPAID_PER_IP", 2)
    db.raffles.update_one({"slug": raffle["slug"]}, {"$set": {"max_tickets_per_person": 0}})
    assert claim(client, raffle["slug"], 1, phone="5500000001").status_code == 200
    assert claim(client, raffle["slug"], 2, phone="5500000002").status_code == 200
    assert claim(client, raffle["slug"], 3, phone="5500000003").status_code == 429


def test_rate_limit_is_per_ip_not_global(client, raffle):
    s = raffle["slug"]
    for i in range(10):
        client.post(f"/api/v1/public/raffles/{s}/claim", json={"folio": 99, "name": "Ana", "phone": PHONE},
                    headers={"x-forwarded-for": "1.1.1.1, 9.9.9.9"})
    blocked = client.post(f"/api/v1/public/raffles/{s}/claim", json={"folio": 99, "name": "Ana", "phone": PHONE},
                          headers={"x-forwarded-for": "1.1.1.1, 9.9.9.9"})
    other = client.post(f"/api/v1/public/raffles/{s}/claim", json={"folio": 99, "name": "Ana", "phone": PHONE},
                        headers={"x-forwarded-for": "2.2.2.2, 9.9.9.9"})
    assert blocked.status_code == 429
    assert other.status_code != 429  # otra IP no se ve afectada


def test_login_rate_limit(client):
    for _ in range(8):
        client.post("/api/v1/auth/login", json={"phone": "5500000000", "password": "mala-mala"})
    assert client.post("/api/v1/auth/login", json={"phone": "5500000000", "password": "mala-mala"}).status_code == 429


# ---------- liberar ----------
def test_public_release_needs_code(client, raffle):
    s = raffle["slug"]
    code = claim(client, s, 1).json()["code"]
    url = f"/api/v1/public/raffles/{s}/release"
    assert client.post(url, json={"folio": 1, "phone": PHONE, "code": "ZZZZ"}).status_code == 403
    assert client.post(url, json={"folio": 1, "phone": "5500000099", "code": code}).status_code == 403
    assert client.post(url, json={"folio": 1, "phone": PHONE, "code": code}).status_code == 200
    assert claim(client, s, 1, phone="5599887766").status_code == 200  # volvió a estar libre


def test_public_release_blocked_after_scratch(client, raffle):
    s = raffle["slug"]
    code = claim(client, s, 1).json()["code"]
    assert client.post("/api/v1/public/scratch", json={"folio": 1, "code": code, "raffle_slug": s}).status_code == 200
    r = client.post(f"/api/v1/public/raffles/{s}/release", json={"folio": 1, "phone": PHONE, "code": code})
    assert r.status_code == 400


# ---------- pagos ----------
def test_cannot_pay_ticket_without_participant(client, admin, raffle):
    client.post(f"/api/v1/raffles/{raffle['id']}/tickets/1/assign", headers=admin, json={})  # entregado sin dueño
    assert pay(client, admin, raffle["id"], 1).status_code == 400


def test_pay_unpay_and_audit(client, admin, db, raffle):
    claim(client, raffle["slug"], 1)
    assert pay(client, admin, raffle["id"], 1).status_code == 200
    r = client.post(f"/api/v1/raffles/{raffle['id']}/tickets/1/unpay", headers=admin, json={"reason": "error"})
    assert r.status_code == 200 and r.json()["status"] == "registered"
    assert {a["action"] for a in db.audit_log.find()} >= {"pay", "unpay"}


def test_assign_does_not_overwrite_other_person_or_downgrade(client, admin, raffle):
    claim(client, raffle["slug"], 1)
    base = f"/api/v1/raffles/{raffle['id']}/tickets/1/assign"
    r = client.post(base, headers=admin, json={"name": "Otra Persona", "phone": "5500000077"})
    assert r.status_code == 400
    r = client.post(base, headers=admin, json={})
    assert r.json()["status"] == "registered"  # no se degrada a 'delivered'


# ---------- sorteo ----------
def test_draw_only_paid_and_only_once(client, admin, raffle):
    rid = raffle["id"]
    assert client.post(f"/api/v1/raffles/{rid}/draw", headers=admin).status_code == 400  # nadie pagó
    claim(client, raffle["slug"], 1)
    claim(client, raffle["slug"], 2, phone="5599887766")
    pay(client, admin, rid, 1)
    r = client.post(f"/api/v1/raffles/{rid}/draw", headers=admin)
    assert r.status_code == 200 and r.json()["winner"]["folio"] == 1
    assert client.post(f"/api/v1/raffles/{rid}/draw", headers=admin).status_code == 400
    info = client.get(f"/api/v1/raffles/{rid}", headers=admin).json()
    assert info["draw_audit"]["eligible_folios"] == [1]


def test_draw_race_has_single_winner(db, admin, client, raffle):
    """Si otro proceso sortea entre la lectura y la escritura, el segundo falla."""
    rid = raffle["id"]
    claim(client, raffle["slug"], 1)
    pay(client, admin, rid, 1)
    real = rs._raffle_of
    def stale(db_, oid_):
        doc = real(db_, oid_)
        db_.raffles.update_one({"_id": oid_}, {"$set": {"status": "drawn", "winner": {"folio": 7}}})  # el "otro" gana
        return doc
    rs._raffle_of = stale
    try:
        with pytest.raises(ValueError):
            rs.run_draw(db, rid)
    finally:
        rs._raffle_of = real
    assert db.raffles.find_one({"_id": ObjectId(rid)})["winner"]["folio"] == 7  # no fue pisado


def test_no_changes_after_draw(client, admin, raffle):
    rid = raffle["id"]
    claim(client, raffle["slug"], 1)
    claim(client, raffle["slug"], 2, phone="5599887766")
    pay(client, admin, rid, 1)
    client.post(f"/api/v1/raffles/{rid}/draw", headers=admin)
    assert pay(client, admin, rid, 2).status_code == 400
    assert client.post(f"/api/v1/raffles/{rid}/tickets/2/release", headers=admin).status_code == 400


def test_delete_blocked_with_payments(client, admin, raffle):
    claim(client, raffle["slug"], 1)
    pay(client, admin, raffle["id"], 1)
    assert client.delete(f"/api/v1/raffles/{raffle['id']}", headers=admin).status_code == 400


# ---------- mantenimiento ----------
def test_auto_release_and_close(db, client, raffle):
    s = raffle["slug"]
    claim(client, s, 1)
    claim(client, s, 2, phone="5599887766")
    old = datetime.now(timezone.utc) - timedelta(hours=100)
    db.tickets.update_many({"folio": {"$in": [1, 2]}}, {"$set": {"registered_at": old}})
    db.tickets.update_one({"folio": 2}, {"$set": {"status": "paid"}})
    run_maintenance(db)
    assert db.tickets.find_one({"folio": 1})["status"] == "free"
    assert db.tickets.find_one({"folio": 2})["status"] == "paid"  # un pago nunca se libera
    assert db.audit_log.count_documents({"action": "auto_release"}) == 1


def test_draw_date_validation_and_local_close(client, admin, db):
    base = {"title": "Rifa X", "prize": "Premio", "prize_value": 1, "price_min": 1, "price_max": 3}
    assert client.post("/api/v1/raffles", headers=admin, json={**base, "draw_date": "10/10/2026"}).status_code == 422
    tomorrow = (datetime.now() + timedelta(days=2)).strftime("%Y-%m-%d")
    r = client.post("/api/v1/raffles", headers=admin, json={**base, "draw_date": tomorrow})
    assert r.status_code == 201
    run_maintenance(db)
    assert db.raffles.find_one({"slug": r.json()["slug"]})["status"] == "open"  # aún no vence


# ---------- Excel ----------
def test_excel_formula_is_text(client, admin, raffle):
    from io import BytesIO
    from openpyxl import load_workbook
    claim(client, raffle["slug"], 1, name='=HYPERLINK("http://evil","x")')
    r = client.get(f"/api/v1/raffles/{raffle['id']}/export/participants", headers=admin)
    ws = load_workbook(BytesIO(r.content))["Participantes"]
    cell = next(c for row in ws.iter_rows() for c in row if isinstance(c.value, str) and "HYPERLINK" in c.value)
    assert cell.data_type != "f"


def test_upload_rejects_fake_image(client, admin, monkeypatch):
    monkeypatch.setenv("CLOUDINARY_CLOUD_NAME", "x")
    from app.config import get_settings
    get_settings.cache_clear()
    r = client.post("/api/v1/upload-image", headers=admin, files={"file": ("a.png", b"not an image", "image/png")})
    get_settings.cache_clear()
    assert r.status_code == 400


# ---------- sincronización del navegador del participante ----------
def test_mine_detects_ticket_released_by_staff(client, admin, raffle):
    s = raffle["slug"]
    code = claim(client, s, 1).json()["code"]
    url = f"/api/v1/public/raffles/{s}/mine"
    body = {"phone": PHONE, "tickets": [{"folio": 1, "code": code}]}
    assert client.post(url, json=body).json()["tickets"][0]["valid"] is True
    client.post(f"/api/v1/raffles/{raffle['id']}/tickets/1/release", headers=admin)
    assert client.post(url, json=body).json()["tickets"][0]["valid"] is False  # código rotado
    # y aunque otra persona lo tome, el código viejo no le sirve a la anterior
    claim(client, s, 1, phone="5599887766")
    assert client.post(url, json=body).json()["tickets"][0]["valid"] is False


# ---------- sesiones ----------
def test_password_change_revokes_old_sessions(client, admin):
    old_refresh = client.post("/api/v1/auth/login", json={"phone": "5500000000", "password": "TestPassword123!"}).json()["refresh_token"]
    r = client.post("/api/v1/change-password", headers=admin,
                    json={"current_password": "TestPassword123!", "new_password": "OtraClave456!"})
    assert r.status_code == 200
    assert client.get("/api/v1/me", headers=admin).status_code == 401  # token viejo muerto
    assert client.post("/api/v1/auth/refresh", json={"refresh_token": old_refresh}).status_code == 401
    fresh = {"Authorization": f"Bearer {r.json()['access_token']}"}
    assert client.get("/api/v1/me", headers=fresh).status_code == 200  # el de quien cambió sigue vivo


def test_deactivated_user_loses_access_immediately(client, admin, db):
    db.users.update_one({"phone": "5500000000"}, {"$set": {"active": False}})
    assert client.get("/api/v1/me", headers=admin).status_code == 401


# =====================================================================
# LOTE 1: robustez (códigos, teléfonos, auto-liberación, sorteo, respaldo)
# =====================================================================
import app.routers.public as pub


def access(client, slug, folio, code, ip="1.1.1.1"):
    return client.post("/api/v1/public/access", json={"raffle_slug": slug, "folio": folio, "code": code},
                       headers={"x-forwarded-for": f"{ip}, 9.9.9.9"})


def test_codes_are_six_chars(db, raffle):
    assert {len(t["access_code"]) for t in db.tickets.find()} == {6}


def test_mine_is_not_a_code_oracle(client, raffle):
    """Con el código correcto pero sin el teléfono del titular, /mine no confirma nada."""
    s = raffle["slug"]
    code = claim(client, s, 1).json()["code"]
    url = f"/api/v1/public/raffles/{s}/mine"
    wrong_phone = client.post(url, json={"phone": "5500000099", "tickets": [{"folio": 1, "code": code}]})
    assert wrong_phone.json()["tickets"][0]["valid"] is False
    assert client.post(url, json={"tickets": [{"folio": 1, "code": code}]}).status_code == 422  # teléfono obligatorio


def test_mine_cost_counts_every_ticket(client, raffle):
    """20 códigos por petición ya no multiplican el ritmo de intentos: cada uno cuenta."""
    url = f"/api/v1/public/raffles/{raffle['slug']}/mine"
    body = {"phone": PHONE, "tickets": [{"folio": i, "code": "AAAAAA"} for i in range(1, 21)]}
    codes = [client.post(url, json=body).status_code for _ in range(5)]
    assert codes[:3] == [200, 200, 200] and 429 in codes  # 60 intentos/min = 3 peticiones de 20


def test_ticket_locks_after_ten_wrong_codes(client, db, raffle):
    s = raffle["slug"]
    code = claim(client, s, 1).json()["code"]
    for i in range(10):
        assert access(client, s, 1, "ZZZZZZ", ip=f"2.2.2.{i}").status_code == 401
    r = access(client, s, 1, code, ip="2.2.2.99")  # ni siquiera el correcto pasa durante el bloqueo
    assert r.status_code == 401 and "Demasiados intentos" in r.json()["detail"]
    db.tickets.update_one({"folio": 1}, {"$set": {"locked_until": datetime.now(timezone.utc) - timedelta(minutes=1)}})
    assert access(client, s, 1, code, ip="2.2.2.98").status_code == 200  # vence el bloqueo


def test_unknown_folio_and_wrong_code_look_identical(client, raffle):
    a = access(client, raffle["slug"], 999, "AAAAAA")
    b = access(client, raffle["slug"], 1, "AAAAAA", ip="3.3.3.3")
    assert a.json()["detail"] == b.json()["detail"]


def test_failure_budget_is_per_raffle_not_per_ip(client, raffle, monkeypatch):
    """Cambiando de IP (X-Forwarded-For falso) el atacante igual se topa con el techo del sorteo."""
    monkeypatch.setattr(pub, "CODE_FAIL_BUDGET", 5)
    s = raffle["slug"]
    codes = [access(client, s, 999, "AAAAAA", ip=f"7.7.7.{i}").status_code for i in range(8)]
    assert codes[:5] == [401] * 5 and set(codes[5:]) == {429}


def test_login_failures_capped_per_phone_even_with_rotating_ips(client):
    codes = []
    for i in range(22):
        r = client.post("/api/v1/auth/login", json={"phone": "5500000000", "password": "mala-mala-1"},
                        headers={"x-forwarded-for": f"8.8.8.{i}, 9.9.9.9"})
        codes.append(r.status_code)
    assert codes[:20] == [401] * 20 and codes[20:] == [429, 429]


def test_phone_prefixes_count_as_same_person(client, raffle):
    s = raffle["slug"]  # límite 2
    assert claim(client, s, 1, phone="5511112222").status_code == 200
    assert claim(client, s, 2, phone="525511112222").status_code == 200
    assert claim(client, s, 3, phone="+52 1 55 1111 2222").status_code == 400


def test_claim_limit_race_is_rolled_back(client, db, raffle, monkeypatch):
    s = raffle["slug"]
    db.raffles.update_one({"slug": s}, {"$set": {"max_tickets_per_person": 1}})
    assert claim(client, s, 1).status_code == 200
    real, calls = rs.count_tickets_of_phone, {"n": 0}
    def racy(*a, **k):  # el primer conteo (previo) "no ve" la otra petición simultánea
        calls["n"] += 1
        return 0 if calls["n"] == 1 else real(*a, **k)
    monkeypatch.setattr(rs, "count_tickets_of_phone", racy)
    assert claim(client, s, 2).status_code == 400
    assert db.tickets.find_one({"folio": 2})["status"] == "free"  # se deshizo


def test_claim_retry_after_lost_response_returns_same_code(client, db, raffle):
    s = raffle["slug"]
    hdr = {"x-forwarded-for": "5.5.5.5, 9.9.9.9"}
    url = f"/api/v1/public/raffles/{s}/claim"
    body = {"folio": 1, "name": "Ana Pérez", "phone": PHONE}
    first = client.post(url, json=body, headers=hdr).json()
    again = client.post(url, json=body, headers=hdr)
    assert again.status_code == 200 and again.json()["code"] == first["code"] and again.json()["resumed"]
    other_net = client.post(url, json=body, headers={"x-forwarded-for": "6.6.6.6, 9.9.9.9"})
    assert other_net.status_code == 400  # desde otra red no se entrega el código
    other_phone = client.post(url, json={**body, "phone": "5599887766"}, headers=hdr)
    assert other_phone.status_code == 400
    db.tickets.update_one({"folio": 1}, {"$set": {"status": "scratched"}})
    assert client.post(url, json=body, headers=hdr).status_code == 400  # ya raspado: no se reentrega


def test_public_board_exposes_ticket_limit(client, raffle):
    b = client.get(f"/api/v1/public/raffles/{raffle['slug']}/board").json()
    assert b["max_tickets_per_person"] == 2
    assert client.get(f"/api/v1/public/raffles/{raffle['slug']}").json()["max_tickets_per_person"] == 2


# ---------- auto-liberación ----------
def _age(db, folio, hours=100):
    db.tickets.update_one({"folio": folio}, {"$set": {"registered_at": datetime.now(timezone.utc) - timedelta(hours=hours)}})


def test_auto_release_never_touches_closed_raffles(client, admin, db, raffle):
    s = raffle["slug"]
    claim(client, s, 1)
    _age(db, 1)
    client.post(f"/api/v1/raffles/{raffle['id']}/close", headers=admin)
    run_maintenance(db)
    assert db.tickets.find_one({"folio": 1})["status"] == "registered"


def test_paid_report_pauses_auto_release(client, db, raffle):
    s = raffle["slug"]
    code = claim(client, s, 1).json()["code"]
    claim(client, s, 2, phone="5599887766")
    r = client.post(f"/api/v1/public/raffles/{s}/paid-report", json={"folio": 1, "code": code})
    assert r.status_code == 200
    bad = client.post(f"/api/v1/public/raffles/{s}/paid-report", json={"folio": 1, "code": "ZZZZZZ"})
    assert bad.status_code == 400
    _age(db, 1)
    _age(db, 2)
    run_maintenance(db)
    assert db.tickets.find_one({"folio": 1})["status"] == "registered"  # avisó: se conserva
    assert db.tickets.find_one({"folio": 2})["status"] == "free"        # no avisó: se libera


# ---------- sorteo: candado y carreras ----------
def _two_paid(client, admin, raffle):
    claim(client, raffle["slug"], 1)
    claim(client, raffle["slug"], 2, phone="5599887766")
    pay(client, admin, raffle["id"], 1)
    pay(client, admin, raffle["id"], 2)


def test_payments_blocked_while_drawing(client, admin, db, raffle):
    claim(client, raffle["slug"], 1)
    db.raffles.update_one({"slug": raffle["slug"]}, {"$set": {"status": "drawing"}})
    assert pay(client, admin, raffle["id"], 1).status_code == 400


def test_failed_draw_releases_the_lock(client, admin, db, raffle):
    r = client.post(f"/api/v1/raffles/{raffle['id']}/draw", headers=admin)  # nadie pagó
    assert r.status_code == 400
    assert db.raffles.find_one({"slug": raffle["slug"]})["status"] == "open"
    assert client.post(f"/api/v1/raffles/{raffle['id']}/close", headers=admin).status_code == 200


def test_winner_unpaid_mid_draw_is_never_selected(client, admin, db, raffle, monkeypatch):
    """Si el 'elegido' deja de estar pagado justo al elegirlo, se reintenta con los pagados reales."""
    _two_paid(client, admin, raffle)
    real_choice, state = rs.secrets.choice, {"n": 0}
    def sneaky(seq):
        pick = real_choice(seq)
        state["n"] += 1
        if state["n"] == 1:  # alguien deshace el pago del elegido antes de sellarlo
            db.tickets.update_one({"_id": pick["_id"]}, {"$set": {"status": "registered"}})
        return pick
    monkeypatch.setattr(rs.secrets, "choice", sneaky)
    r = client.post(f"/api/v1/raffles/{raffle['id']}/draw", headers=admin)
    assert r.status_code == 200
    winner = db.tickets.find_one({"folio": r.json()["winner"]["folio"]})
    assert winner["status"] == "paid" and winner["is_winner"] is True
    assert db.tickets.count_documents({"is_winner": True}) == 1


def test_draw_fails_cleanly_if_every_payment_vanishes(client, admin, db, raffle, monkeypatch):
    _two_paid(client, admin, raffle)
    real_choice = rs.secrets.choice
    def vanish(seq):
        pick = real_choice(seq)
        db.tickets.update_many({}, {"$set": {"status": "registered"}})
        return pick
    monkeypatch.setattr(rs.secrets, "choice", vanish)
    r = client.post(f"/api/v1/raffles/{raffle['id']}/draw", headers=admin)
    assert r.status_code == 400
    assert db.raffles.find_one({"slug": raffle["slug"]})["status"] == "open"
    assert db.tickets.count_documents({"is_winner": True}) == 0


def test_payment_slipping_in_after_seal_is_reverted(client, admin, db, raffle, monkeypatch):
    """Un pago que se cuela justo cuando el sorteo ya tiene su lista queda deshecho."""
    claim(client, raffle["slug"], 1)
    claim(client, raffle["slug"], 2, phone="5599887766")
    pay(client, admin, raffle["id"], 1)
    client.post(f"/api/v1/raffles/{raffle['id']}/draw", headers=admin)
    # se salta el chequeo previo para simular la ventana de carrera
    monkeypatch.setattr(rs, "_ensure_not_drawn", lambda raffle: None)
    r = pay(client, admin, raffle["id"], 2)
    assert r.status_code == 400
    assert db.tickets.find_one({"folio": 2})["status"] == "registered"


def test_stuck_drawing_is_unlocked_by_maintenance(db, client, raffle):
    db.raffles.update_one({"slug": raffle["slug"]}, {"$set": {
        "status": "drawing", "prev_status": "open",
        "drawing_at": datetime.now(timezone.utc) - timedelta(minutes=30)}})
    run_maintenance(db)
    assert db.raffles.find_one({"slug": raffle["slug"]})["status"] == "open"


# ---------- respaldo ----------
def test_backup_works_after_editing_the_raffle(client, admin, raffle):
    claim(client, raffle["slug"], 1)
    pay(client, admin, raffle["id"], 1)
    assert client.patch(f"/api/v1/raffles/{raffle['id']}", headers=admin, json={"notes": "editado"}).status_code == 200
    r = client.get(f"/api/v1/raffles/{raffle['id']}/export/backup", headers=admin)
    assert r.status_code == 200
    data = r.json()
    assert data["backup_version"] == 2 and data["raffle"]["notes"] == "editado"
    assert any(a["action"] == "pay" for a in data["audit_log"])


def test_staff_assign_normalizes_phone_for_the_limit(client, admin, raffle):
    base = f"/api/v1/raffles/{raffle['id']}/tickets"
    assert client.post(f"{base}/1/assign", headers=admin, json={"name": "Ana Pérez", "phone": "5511112222"}).status_code == 200
    assert client.post(f"{base}/2/assign", headers=admin, json={"name": "Ana Pérez", "phone": "+52 55 1111 2222"}).status_code == 200
    r = client.post(f"{base}/3/assign", headers=admin, json={"name": "Ana Pérez", "phone": "525511112222"})
    assert r.status_code == 400  # límite 2: es la misma persona
