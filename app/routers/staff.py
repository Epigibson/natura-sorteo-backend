"""Rutas del dashboard (staff): sorteos, boletos, sorteo, exportación."""
from __future__ import annotations

import io
import json
from datetime import datetime, timezone
from typing import Annotated, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse

from app.db import get_db
from app.schemas import (
    ChangePasswordIn,
    RaffleCreate,
    RaffleUpdate,
    RaffleOut,
    RaffleStats,
    TicketAssign,
    TicketOut,
    TicketPaidIn,
    TicketReleaseIn,
    TicketUnpayIn,
)
from app.security import (
    CurrentUser,
    create_access_token,
    create_refresh_token,
    require_admin,
    require_staff,
)
from app.services import raffle_service as rs

router = APIRouter(prefix="/api/v1", tags=["staff"])


@router.get("/me")
def me(user: Annotated[CurrentUser, Depends(require_staff)]) -> dict:
    return {"sub": user["sub"], "role": user["role"], "name": user["name"]}


@router.post("/change-password")
def change_password(
    body: ChangePasswordIn,
    user: Annotated[CurrentUser, Depends(require_staff)],
) -> dict:
    """Cambia la contraseña del usuario autenticado."""
    from bson import ObjectId

    from app.security import hash_password, verify_password

    db = get_db()
    try:
        doc = db.users.find_one({"_id": ObjectId(user["sub"])})
    except Exception:
        doc = None
    if not doc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Usuario no encontrado")

    if not verify_password(body.current_password, doc.get("password_hash", "")):
        raise HTTPException(status.HTTP_400_BAD_REQUEST, "La contraseña actual es incorrecta")

    if verify_password(body.new_password, doc.get("password_hash", "")):
        raise HTTPException(
            status.HTTP_400_BAD_REQUEST,
            "La nueva contraseña debe ser diferente a la actual",
        )

    # $inc de token_version invalida todas las sesiones anteriores (otros dispositivos, tokens robados)
    new = db.users.find_one_and_update(
        {"_id": doc["_id"]},
        {"$set": {"password_hash": hash_password(body.new_password)}, "$inc": {"token_version": 1}},
        return_document=True,
    )
    tv = new.get("token_version", 0)
    rs.audit(db, "change_password", by=user["sub"])
    # Tokens nuevos para que quien cambia la contraseña no pierda su propia sesión
    return {
        "ok": True,
        "message": "Contraseña actualizada correctamente",
        "access_token": create_access_token(sub=user["sub"], role=new.get("role", "collab"), name=new.get("name", ""), tv=tv),
        "refresh_token": create_refresh_token(sub=user["sub"], role=new.get("role", "collab"), tv=tv),
    }


# ---------- Sorteos ----------
@router.get("/raffles", response_model=list[RaffleOut])
def list_raffles(user: Annotated[CurrentUser, Depends(require_staff)]) -> list[dict]:
    return rs.list_raffles(get_db())


@router.post("/raffles", response_model=RaffleOut, status_code=status.HTTP_201_CREATED)
def create_raffle(
    body: RaffleCreate,
    user: Annotated[CurrentUser, Depends(require_admin)],
) -> dict:
    try:
        return rs.crear_sorteo(
            get_db(),
            title=body.title,
            prize=body.prize,
            prize_value=body.prize_value,
            price_min=body.price_min,
            price_max=body.price_max,
            draw_date=body.draw_date,
            notes=body.notes,
            meet_url=body.meet_url,
            max_tickets_per_person=body.max_tickets_per_person,
        )
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))


@router.get("/raffles/{raffle_id}", response_model=RaffleOut)
def get_raffle(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_staff)],
) -> dict:
    doc = rs.get_raffle(get_db(), raffle_id)
    if not doc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sorteo no encontrado")
    return doc


@router.get("/raffles/{raffle_id}/stats", response_model=RaffleStats)
def raffle_stats(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_staff)],
) -> dict:
    if not rs.get_raffle(get_db(), raffle_id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sorteo no encontrado")
    return rs.stats_sorteo(get_db(), raffle_id)


@router.get("/raffles/{raffle_id}/tickets", response_model=list[TicketOut])
def list_tickets(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_staff)],
    status_filter: Annotated[Optional[str], Query(alias="status")] = None,
) -> list[dict]:
    if not rs.get_raffle(get_db(), raffle_id):
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sorteo no encontrado")
    return rs.list_tickets(get_db(), raffle_id, status_filter)


@router.post("/raffles/{raffle_id}/tickets/{folio}/assign", response_model=TicketOut)
def assign_ticket(
    raffle_id: str,
    folio: int,
    body: TicketAssign,
    user: Annotated[CurrentUser, Depends(require_staff)],
) -> dict:
    try:
        doc = rs.entregar_boleto(get_db(), raffle_id, folio, body.name, body.phone)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    if not doc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Boleto no encontrado")
    return doc


@router.post("/raffles/{raffle_id}/tickets/{folio}/pay", response_model=TicketOut)
def mark_paid(
    raffle_id: str,
    folio: int,
    body: TicketPaidIn,
    user: Annotated[CurrentUser, Depends(require_staff)],
) -> dict:
    try:
        doc = rs.marcar_pagado(get_db(), raffle_id, folio, body.note, by=user["sub"], expected_phone=body.expected_phone)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    if not doc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Boleto no encontrado")
    return doc


@router.post("/raffles/{raffle_id}/tickets/{folio}/unpay", response_model=TicketOut)
def unmark_paid(
    raffle_id: str,
    folio: int,
    body: TicketUnpayIn,
    user: Annotated[CurrentUser, Depends(require_admin)],
) -> dict:
    """Deshace un pago marcado por error (solo admin y antes del sorteo)."""
    try:
        doc = rs.desmarcar_pago(get_db(), raffle_id, folio, by=user["sub"], reason=body.reason)
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    if not doc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Boleto no encontrado")
    return doc


@router.post("/raffles/{raffle_id}/tickets/{folio}/release", response_model=TicketOut)
def release_ticket(
    raffle_id: str,
    folio: int,
    user: Annotated[CurrentUser, Depends(require_staff)],
    body: Optional[TicketReleaseIn] = None,
) -> dict:
    try:
        doc = rs.liberar_boleto(
            get_db(), raffle_id, folio, by=user["sub"],
            expected_phone=body.expected_phone if body else None,
        )
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))
    if not doc:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Boleto no encontrado")
    return doc


@router.post("/raffles/{raffle_id}/close", response_model=RaffleOut)
def close_raffle(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_admin)],
) -> dict:
    try:
        return rs.close_raffle(get_db(), raffle_id, by=user["sub"])
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))


@router.post("/raffles/{raffle_id}/draw")
def draw(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_admin)],
) -> dict:
    try:
        return rs.run_draw(get_db(), raffle_id, by=user["sub"])
    except ValueError as exc:
        raise HTTPException(status.HTTP_400_BAD_REQUEST, str(exc))


# ---------- Exportación Excel ----------
def _text_cell(cell) -> None:
    """Evita inyección de fórmulas: openpyxl trata '=...' como fórmula; lo forzamos a texto."""
    if isinstance(cell.value, str) and cell.value.startswith("="):
        cell.data_type = "s"


@router.get("/raffles/{raffle_id}/export/participants")
def export_participants(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_staff)],
) -> StreamingResponse:
    """Exporta la lista de participantes y boletos a un archivo .xlsx."""
    from datetime import datetime, timezone

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    db = get_db()
    raffle = rs.get_raffle(db, raffle_id)
    if not raffle:
        raise HTTPException(status.HTTP_404_NOT_FOUND, "Sorteo no encontrado")

    tickets = rs.list_tickets(db, raffle_id)
    stats = rs.stats_sorteo(db, raffle_id)

    wb = Workbook()

    # ---- Hoja 1: Participantes ----
    ws = wb.active
    ws.title = "Participantes"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill("solid", fgColor="1B5E20")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    money_fmt = '"$"#,##0'

    # El código de acceso permite raspar/registrar a nombre de otra persona: solo lo ve el admin
    show_codes = user.get("role") == "admin"
    headers = [
        "Folio",
        "Nombre",
        "Teléfono",
        "Monto",
        "Estado",
        "Código" if show_codes else "Código (solo admin)",
        "Registrado",
        "Raspado",
        "Pagado",
        "Ganador",
    ]
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border

    status_labels = {
        "free": "Libre",
        "delivered": "Entregado",
        "registered": "Registrado",
        "scratched": "Raspado",
        "paid": "Pagado",
        "released": "Liberado",
    }
    winner_folio = (raffle.get("winner") or {}).get("folio")

    def _fmt_date(val):
        if not val:
            return ""
        if isinstance(val, str):
            return val[:16].replace("T", " ")
        return val.strftime("%d/%m/%Y %H:%M")

    row = 2
    for t in tickets:
        p = t.get("participant") or {}
        is_winner = t.get("folio") == winner_folio and raffle.get("status") == "drawn"
        values = [
            t.get("folio"),
            p.get("name") or "",
            p.get("phone") or "",
            t.get("amount"),
            status_labels.get(t.get("status"), t.get("status")),
            t.get("access_code") if show_codes else "",
            _fmt_date(t.get("registered_at")),
            _fmt_date(t.get("scratched_at")),
            _fmt_date(t.get("paid_at")),
            "🏆 SÍ" if is_winner else "",
        ]
        for col, v in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=v)
            _text_cell(cell)
            cell.border = thin_border
            if col == 4 and isinstance(v, (int, float)):
                cell.number_format = money_fmt
            if is_winner:
                cell.fill = PatternFill("solid", fgColor="FFF8E1")
            if t.get("status") == "paid" and not is_winner:
                cell.fill = PatternFill("solid", fgColor="F1F8F2")
        row += 1

    # Ajuste de anchos
    widths = [8, 28, 16, 10, 14, 10, 18, 18, 18, 10]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:J{row - 1}"

    # ---- Hoja 2: Resumen ----
    ws2 = wb.create_sheet("Resumen")
    ws2.column_dimensions["A"].width = 32
    ws2.column_dimensions["B"].width = 22

    title_cell = ws2.cell(row=1, column=1, value=f"Sorteo: {raffle['title']}")
    title_cell.font = Font(bold=True, size=14, color="1B5E20")

    ws2.cell(row=2, column=1, value=f"Premio: {raffle['prize']}").font = Font(size=11)
    ws2.cell(row=3, column=1, value=f"Valor del premio: ${raffle['prize_value']} MXN").font = Font(size=11)
    ws2.cell(row=4, column=1, value=f"Rango de precios: ${raffle['price_min']} – ${raffle['price_max']}").font = Font(size=11)
    ws2.cell(row=5, column=1, value=f"Estado: {status_labels.get(raffle['status'], raffle['status'])}").font = Font(size=11)
    ws2.cell(row=6, column=1, value=f"Generado: {datetime.now(timezone.utc).strftime('%d/%m/%Y %H:%M')} UTC").font = Font(size=9, italic=True)

    summary = [
        ("Total de boletos", stats["total_tickets"]),
        ("Libres", stats["free"]),
        ("Entregados", stats["delivered"]),
        ("Registrados", stats["registered"]),
        ("Raspados", stats["scratched"]),
        ("Pagados", stats["paid"]),
        ("Participantes únicos", stats["participants"]),
        ("Recaudación esperada", stats["revenue_expected"]),
        ("Recaudación confirmada", stats["revenue_confirmed"]),
    ]
    r = 8
    for label, val in summary:
        c1 = ws2.cell(row=r, column=1, value=label)
        c1.font = Font(bold=True, size=11)
        c2 = ws2.cell(row=r, column=2, value=val)
        c2.font = Font(size=11)
        if "Recaudación" in label:
            c2.number_format = money_fmt
        r += 1

    if raffle.get("status") == "drawn" and raffle.get("winner"):
        w = raffle["winner"]
        ws2.cell(row=r + 1, column=1, value="🏆 GANADOR").font = Font(bold=True, size=12, color="8D6E00")
        ws2.cell(row=r + 2, column=1, value=f"Folio {w.get('folio')}").font = Font(bold=True, size=11)
        pname = (w.get("participant") or {}).get("name") or "—"
        ws2.cell(row=r + 3, column=1, value=f"Nombre: {pname}").font = Font(size=11)
        ws2.cell(row=r + 4, column=1, value=f"Monto del boleto: ${w.get('amount')}").font = Font(size=11)

    # Generar bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    slug = raffle.get("slug", "sorteo")
    filename = f"participantes-{slug}.xlsx"

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


# ---------- Editar sorteo ----------
@router.patch("/raffles/{raffle_id}", response_model=RaffleOut)
def update_raffle(
    raffle_id: str,
    body: RaffleUpdate,
    user: Annotated[CurrentUser, Depends(require_admin)],
) -> dict:
    """Edita un sorteo (solo si no está sorteado)."""
    from bson import ObjectId
    from datetime import datetime, timezone
    from app.services import raffle_service as rs

    db = get_db()
    try:
        doc = db.raffles.find_one({"_id": ObjectId(raffle_id)})
    except Exception:
        raise HTTPException(400, "ID inválido")
    if not doc:
        raise HTTPException(404, "Sorteo no encontrado")
    if doc.get("status") == "drawn":
        raise HTTPException(400, "No se puede editar un sorteo ya sorteado")

    allowed = {}
    sent = body.model_fields_set  # distingue "no enviado" de "enviado vacío (null)"
    clearable = ("draw_date", "notes", "image_url")
    for field in ("title", "prize", "prize_value", "draw_date", "notes", "image_url", "meet_url", "max_tickets_per_person"):
        if field not in sent:
            continue
        val = getattr(body, field, None)
        if val is None:
            if field in clearable:
                allowed[field] = None  # el usuario vació el campo: se borra de verdad
            continue
        if field == "meet_url":
            val = rs.normalize_meet_url(val)
        allowed[field] = val
    if not allowed:
        raise HTTPException(400, "Nada que actualizar")
    allowed["updated_at"] = datetime.now(timezone.utc)

    db.raffles.update_one({"_id": doc["_id"]}, {"$set": allowed})
    result = rs.get_raffle(db, raffle_id)
    return result


@router.post("/raffles/{raffle_id}/reopen", response_model=RaffleOut)
def reopen_raffle(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_admin)],
) -> dict:
    """Reabre la venta de un sorteo cerrado (cierre por error o fecha mal puesta)."""
    from bson import ObjectId

    db = get_db()
    try:
        doc = db.raffles.find_one({"_id": ObjectId(raffle_id)})
    except Exception:
        raise HTTPException(400, "ID inválido")
    if not doc:
        raise HTTPException(404, "Sorteo no encontrado")
    if doc.get("status") != "closed":
        raise HTTPException(400, "Solo se puede reabrir un sorteo cerrado")
    # Si la fecha ya pasó, el cierre automático lo volvería a cerrar en minutos
    if doc.get("draw_date") and doc["draw_date"] < rs.local_today():
        raise HTTPException(400, "La fecha del sorteo ya pasó; cámbiala por hoy o una futura antes de reabrir la venta")
    res = db.raffles.update_one({"_id": doc["_id"], "status": "closed"}, {"$set": {"status": "open"}})
    if res.modified_count == 0:
        raise HTTPException(400, "El sorteo cambió de estado; recarga e inténtalo de nuevo")
    rs.audit(db, "reopen", by=user["sub"], raffle_id=doc["_id"])
    return rs.get_raffle(db, raffle_id)


# ---------- Historial global de participantes ----------
@router.get("/participants")
def list_participants(
    user: Annotated[CurrentUser, Depends(require_admin)],  # teléfonos de todos los sorteos
) -> dict:
    """Historial de participantes en todos los sorteos."""
    from app.services import raffle_service as rs

    db = get_db()
    raffles = {str(r["_id"]): r for r in db.raffles.find()}
    seen: dict[str, dict] = {}

    for t in db.tickets.find({"participant.phone": {"$ne": None}}):
        p = t.get("participant") or {}
        phone = p.get("phone")
        if not phone:
            continue
        rid = str(t["raffle_id"])
        raffle = raffles.get(rid, {})
        entry = seen.setdefault(phone, {
            "name": p.get("name"),
            "phone": phone,
            "raffles": [],
            "total_spent": 0,
            "tickets_count": 0,
        })
        entry["raffles"].append({
            "raffle": raffle.get("title", "?"),
            "raffle_id": rid,
            "folio": t["folio"],
            "amount": t.get("amount"),
            "status": t.get("status"),
            "is_winner": t.get("is_winner", False),
        })
        entry["tickets_count"] += 1
        if t.get("status") == "paid" and t.get("amount"):
            entry["total_spent"] += t["amount"]

    result = sorted(seen.values(), key=lambda x: x["total_spent"], reverse=True)
    return {"participants": result, "total": len(result)}


# ---------- Backup JSON ----------
@router.get("/raffles/{raffle_id}/export/backup")
def export_backup(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_admin)],  # incluye códigos de acceso
) -> dict:
    """Exporta todo el sorteo en JSON (sorteo + boletos + participantes)."""
    import json
    from fastapi.responses import StreamingResponse
    from app.services import raffle_service as rs

    db = get_db()
    raffle = rs.get_raffle(db, raffle_id)
    if not raffle:
        raise HTTPException(404, "Sorteo no encontrado")
    tickets = rs.list_tickets(db, raffle_id)
    stats = rs.stats_sorteo(db, raffle_id)

    # Bitácora del sorteo (pagos, liberaciones, sorteo) para poder reconstruir lo ocurrido
    audit_rows = [
        {k: v for k, v in row.items() if k != "_id"}
        for row in db.audit_log.find({"raffle_id": raffle_id}).sort("at", 1)
    ]

    data = {
        "backup_version": 2,
        "exported_at": datetime.now(timezone.utc),
        "raffle": raffle,
        "stats": stats,
        "tickets": tickets,
        "audit_log": audit_rows,
    }

    # Cualquier fecha/ObjectId se serializa solo: antes un campo nuevo (p. ej. updated_at tras
    # editar) tumbaba todo el respaldo con un 500.
    def _default(obj):
        if hasattr(obj, "isoformat"):
            return obj.isoformat()
        return str(obj)

    json_str = json.dumps(data, ensure_ascii=False, indent=2, default=_default)
    buf = io.BytesIO(json_str.encode("utf-8"))
    return StreamingResponse(
        buf,
        media_type="application/json",
        headers={
            "Content-Disposition": f"attachment; filename=backup-{raffle.get('slug', 'sorteo')}.json"
        },
    )

# ---------- Upload imagen a Cloudinary ----------
@router.post("/upload-image")
async def upload_image(
    file: UploadFile,
    user: Annotated[CurrentUser, Depends(require_staff)],
) -> dict:
    """Sube una imagen a Cloudinary y devuelve la URL."""
    import cloudinary
    import cloudinary.uploader

    from app.config import get_settings

    s = get_settings()
    if not s.cloudinary_cloud_name:
        raise HTTPException(500, "Cloudinary no está configurado")

    cloudinary.config(
        cloud_name=s.cloudinary_cloud_name,
        api_key=s.cloudinary_api_key,
        api_secret=s.cloudinary_api_secret,
        secure=True,
    )

    # Validar tipo de archivo
    allowed_types = {"image/jpeg", "image/png", "image/webp", "image/gif", "image/jpg"}
    if file.content_type not in allowed_types:
        raise HTTPException(400, "Solo se permiten imágenes JPG, PNG, WebP o GIF")

    content = await file.read(5 * 1024 * 1024 + 1)
    if len(content) > 5 * 1024 * 1024:
        raise HTTPException(400, "Imagen muy grande (máx 5 MB)")
    # Verificar la firma real del archivo, no solo el content-type declarado
    if not (content.startswith((b"\xff\xd8\xff", b"\x89PNG\r\n\x1a\n", b"GIF8")) or (content[:4] == b"RIFF" and content[8:12] == b"WEBP")):
        raise HTTPException(400, "El archivo no parece una imagen válida")
    try:
        import base64
        b64 = base64.b64encode(content).decode("utf-8")
        mime = file.content_type or "image/jpeg"
        result = cloudinary.uploader.upload(
            f"data:{mime};base64,{b64}",
            folder="sorteo-natura",
            resource_type="image",
        )
        return {"url": result.get("secure_url"), "public_id": result.get("public_id")}
    except Exception as exc:
        raise HTTPException(500, f"Error al subir imagen: {exc}")

# ---------- Eliminar sorteo (cascade) ----------
@router.delete("/raffles/{raffle_id}")
def delete_raffle(
    raffle_id: str,
    user: Annotated[CurrentUser, Depends(require_admin)],
) -> dict:
    """Elimina un sorteo y todos sus boletos (cascade)."""
    from bson import ObjectId

    db = get_db()
    try:
        rid = ObjectId(raffle_id)
    except Exception:
        raise HTTPException(400, "ID inválido")

    raffle = db.raffles.find_one({"_id": rid})
    if not raffle:
        raise HTTPException(404, "Sorteo no encontrado")

    if raffle.get("status") == "drawn" or db.tickets.count_documents({"raffle_id": rid, "status": "paid"}) > 0:
        raise HTTPException(
            400,
            "No se puede eliminar un sorteo con boletos pagados o ya sorteado. "
            "Descarga el respaldo y pide a soporte borrarlo manualmente si es necesario.",
        )

    # Eliminar boletos primero
    tickets_result = db.tickets.delete_many({"raffle_id": rid})
    # Eliminar sorteo
    db.raffles.delete_one({"_id": rid})
    rs.audit(db, "delete_raffle", by=user["sub"], raffle_id=rid, title=raffle.get("title"), tickets=tickets_result.deleted_count)

    return {
        "ok": True,
        "deleted_raffle": raffle.get("title"),
        "deleted_tickets": tickets_result.deleted_count,
    }
